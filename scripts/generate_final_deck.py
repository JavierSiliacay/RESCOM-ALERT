import os
from PIL import Image
import qrcode

# 1. Prepare QR code
def prepare_qr():
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=12,
        border=2,
    )
    qr.add_data("https://rescom-alert.vercel.app/download")
    qr.make(fit=True)
    img = qr.make_image(fill_color="#065f46", back_color="white").convert("RGB")
    os.makedirs("public", exist_ok=True)
    img.save("public/download_qr.png", quality=95)
    print("Download QR code created at public/download_qr.png")

prepare_qr()

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck_with_download_slide():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1:1 System Palette
    C_BG = RGBColor(248, 250, 252)          # slate-50
    C_CARD = RGBColor(255, 255, 255)        # pure white
    C_BORDER = RGBColor(226, 232, 240)      # slate-200
    C_TEXT = RGBColor(15, 23, 42)           # slate-900
    C_TEXT_BODY = RGBColor(71, 85, 105)     # slate-600
    C_MUTED = RGBColor(148, 163, 184)       # slate-400

    # Accents
    C_EMERALD = RGBColor(6, 95, 70)         # emerald-800
    C_EMERALD_TEXT = RGBColor(4, 120, 87)   # emerald-700
    C_EMERALD_SOFT = RGBColor(236, 253, 245)# emerald-50
    C_EMERALD_BORDER = RGBColor(167, 243, 208)

    C_BLUE = RGBColor(29, 78, 216)          # blue-700
    C_BLUE_SOFT = RGBColor(239, 246, 255)   # blue-50
    C_BLUE_BORDER = RGBColor(191, 219, 254)

    C_AMBER = RGBColor(217, 119, 6)         # amber-600
    C_AMBER_SOFT = RGBColor(255, 251, 235)  # amber-50
    C_AMBER_BORDER = RGBColor(252, 211, 77)

    C_RED = RGBColor(220, 38, 38)           # red-600
    C_RED_SOFT = RGBColor(254, 242, 242)    # red-50

    logo_path = os.path.abspath("public/rescom-emblem.jpg")
    seal_path = os.path.abspath("public/rescom-pa-seal.png")
    bg_slide_img = os.path.abspath("public/rescom-slide-bg.jpg")
    qr_img = os.path.abspath("public/download_qr.png")

    icon_web = os.path.abspath("public/icons_diagram/icon_web.png")
    icon_tower = os.path.abspath("public/icons_diagram/icon_tower.png")
    icon_phone = os.path.abspath("public/icons_diagram/icon_phone.png")
    icon_check = os.path.abspath("public/icons_diagram/icon_check.png")

    blank_layout = prs.slide_layouts[6]

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, tag, title, slide_num, total=7):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.42), Inches(2.8), Inches(0.36))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_EMERALD_SOFT
        badge.line.color.rgb = C_EMERALD_BORDER
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = tag.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = C_EMERALD_TEXT
        p_b.alignment = PP_ALIGN.CENTER

        tb = slide.shapes.add_textbox(Inches(1.0), Inches(0.85), Inches(9.5), Inches(0.65))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(25)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        if os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(11.6), Inches(0.42), width=Inches(0.75))

        tb_f = slide.shapes.add_textbox(Inches(1.0), Inches(6.9), Inches(8.0), Inches(0.3))
        tf_f = tb_f.text_frame
        tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
        p = tf_f.paragraphs[0]
        p.text = "10RCDG RESCOM-ALERT  •  Technical Briefing & Command Defense"
        p.font.size = Pt(10)
        p.font.color.rgb = C_MUTED

        tb_n = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.0), Inches(0.3))
        tf_n = tb_n.text_frame
        tf_n.margin_left = tf_n.margin_top = tf_n.margin_right = tf_n.margin_bottom = 0
        p = tf_n.paragraphs[0]
        p.text = f"{slide_num:02d} / {total:02d}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # SLIDE 1: CLEAN COVER WITH SUBTLE PHOTO BACKGROUND
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    if os.path.exists(bg_slide_img):
        s1.shapes.add_picture(bg_slide_img, 0, 0, width=Inches(13.333), height=Inches(7.5))
    else:
        set_bg(s1)

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(0.85), Inches(10.333), Inches(5.8))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_CARD
    card1.line.color.rgb = C_BORDER
    card1.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(5.5), Inches(1.15), width=Inches(1.05))
    if os.path.exists(seal_path):
        s1.shapes.add_picture(seal_path, Inches(6.8), Inches(1.15), width=Inches(1.05))

    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(2.35), Inches(4.7), Inches(0.36))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = C_EMERALD_SOFT
    badge1.line.color.rgb = C_EMERALD_BORDER
    badge1.line.width = Pt(1)
    tf = badge1.text_frame
    p = tf.paragraphs[0]
    p.text = "PHILIPPINE ARMY RESERVE COMMAND"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT
    p.alignment = PP_ALIGN.CENTER

    tb = s1.shapes.add_textbox(Inches(2.0), Inches(2.85), Inches(9.333), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "10RCDG RESCOM ALERT"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Tactical Mass Emergency Siren & Rapid Mobilization Notification Architecture"
    p2.font.size = Pt(16)
    p2.font.color.rgb = C_TEXT_BODY
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = "10th Regional Community Defense Group  •  Camp General Manuel T Yan Sr, Mawab"
    p3.font.size = Pt(12)
    p3.font.color.rgb = C_MUTED
    p3.alignment = PP_ALIGN.CENTER

    pills = [
        ("Fast Group SMS", "Offline GSM Payload", C_EMERALD, C_EMERALD_SOFT, C_EMERALD_BORDER, Inches(2.2)),
        ("Hardware Siren Override", "Bypasses Silent / Mute", C_AMBER, C_AMBER_SOFT, C_AMBER_BORDER, Inches(5.2)),
        ("Instant Headcount", "Real-Time Telemetry", C_BLUE, C_BLUE_SOFT, C_BLUE_BORDER, Inches(8.2))
    ]
    for p_title, p_sub, p_acc, p_bg, p_bdr, p_x in pills:
        pill_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, p_x, Inches(4.7), Inches(2.9), Inches(0.65))
        pill_box.fill.solid()
        pill_box.fill.fore_color.rgb = p_bg
        pill_box.line.color.rgb = p_bdr
        pill_box.line.width = Pt(1)
        tf_pb = pill_box.text_frame
        p_t = tf_pb.paragraphs[0]
        p_t.text = p_title
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = p_acc
        p_t.alignment = PP_ALIGN.CENTER
        p_s = tf_pb.add_paragraph()
        p_s.text = p_sub
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = C_TEXT_BODY
        p_s.alignment = PP_ALIGN.CENTER

    tb_sub = s1.shapes.add_textbox(Inches(2.0), Inches(5.5), Inches(9.333), Inches(0.9))
    tf_sub = tb_sub.text_frame
    p = tf_sub.paragraphs[0]
    p.text = "Presented by: Javier Siliacay (aka vier from USTP Autotronics)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_TEXT
    p.alignment = PP_ALIGN.CENTER

    p2 = tf_sub.add_paragraph()
    p2.text = "Technical Briefing & Command Defense  •  October 2026"
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_MUTED
    p2.alignment = PP_ALIGN.CENTER

    s1.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "Good day, Colonel, and members of the 10RCDG Command Staff.\n"
        "Today I am presenting RESCOM-ALERT — a purpose-built tactical technology designed to solve one critical problem: ensuring emergency mobilization orders wake up and alert our reservists, even when their phones are locked, muted, or offline."
    )

    # ==========================================
    # SLIDE 2: THE PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "The Problem", "Why Conventional Messaging Fails in Emergencies", 2, 7)

    problems = [
        ("74% SILENT", "Phones on Mute & DND", "Troops sleep through or miss standard SMS dings and chat notifications during Red Alerts.", C_RED, C_RED_SOFT),
        ("NO DATA", "Chat Apps Fail in Storms", "Messenger and Viber require active internet. When cell towers lose data, chat groups fail.", C_AMBER, C_AMBER_SOFT),
        ("NO HEADCOUNT", "Zero Acknowledgment", "Commanders have no way to verify who received the order and who is actively en route.", C_TEXT, RGBColor(241, 245, 249))
    ]

    left = Inches(1.0)
    for stat, header, desc, accent, bg_pill in problems:
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.5), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        b_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.4), Inches(2.3), Inches(2.7), Inches(0.8))
        b_box.fill.solid()
        b_box.fill.fore_color.rgb = bg_pill
        b_box.line.color.rgb = accent
        b_box.line.width = Pt(1)
        tf_b = b_box.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = stat
        p_b.font.size = Pt(20)
        p_b.font.bold = True
        p_b.font.color.rgb = accent
        p_b.alignment = PP_ALIGN.CENTER

        tb = s2.shapes.add_textbox(left + Inches(0.35), Inches(3.4), Inches(2.8), Inches(2.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = header
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_BODY

        left += Inches(3.9)

    # ==========================================
    # SLIDE 3: SYSTEM DESIGN DIAGRAM WITH CLEAN ICONS
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "System Architecture", "System Design: Signal Flow & Hardware Override", 3, 7)

    node_w = Inches(2.4)
    node_h = Inches(2.7)
    y_top = Inches(1.8)

    nodes = [
        ("01. DISPATCH", "Command Web Portal", icon_web, [
            "Officer initiates alert",
            "Targets CDC battalions",
            "Fires Siren Command"
        ], C_EMERALD, C_EMERALD_SOFT, C_EMERALD_BORDER, Inches(1.0)),

        ("02. TRANSMIT", "Telco Airwave", icon_tower, [
            "Smart / Globe network",
            "Basic GSM cellular text",
            "ZERO mobile data req."
        ], C_BLUE, C_BLUE_SOFT, C_BLUE_BORDER, Inches(4.0)),

        ("03. OVERRIDE", "Android Field App", icon_phone, [
            "Bypasses Silent / Mute",
            "Wakes locked screen",
            "Loud siren sounds"
        ], C_AMBER, C_AMBER_SOFT, RGBColor(252, 211, 77), Inches(7.0)),

        ("04. ACKNOWLEDGE", "Troop Compliance", icon_check, [
            "Soldier reads order",
            "Taps '✓ I Received Order'",
            "Siren silences instantly"
        ], C_EMERALD, C_EMERALD_SOFT, C_EMERALD_BORDER, Inches(10.0))
    ]

    for step_tag, subtitle, icon_file, bullets, accent, bg_soft, border_col, x_pos in nodes:
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, y_top, node_w, node_h)
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD
        box.line.color.rgb = border_col
        box.line.width = Pt(1.5)

        if os.path.exists(icon_file):
            s3.shapes.add_picture(icon_file, x_pos + Inches(0.8), y_top + Inches(0.18), width=Inches(0.8))

        tb_t = s3.shapes.add_textbox(x_pos + Inches(0.1), y_top + Inches(1.02), node_w - Inches(0.2), Inches(0.5))
        tf_t = tb_t.text_frame
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = step_tag
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = accent
        p_t.alignment = PP_ALIGN.CENTER

        p_s = tf_t.add_paragraph()
        p_s.text = subtitle
        p_s.font.size = Pt(10.5)
        p_s.font.bold = True
        p_s.font.color.rgb = C_TEXT
        p_s.alignment = PP_ALIGN.CENTER

        tb_b = s3.shapes.add_textbox(x_pos + Inches(0.18), y_top + Inches(1.6), node_w - Inches(0.36), Inches(1.0))
        tf_b = tb_b.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        tf_b.word_wrap = True

        for idx, b in enumerate(bullets):
            p = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
            p.text = "• " + b
            p.font.size = Pt(9)
            p.font.color.rgb = C_TEXT_BODY

    arrow_xs = [Inches(3.45), Inches(6.45), Inches(9.45)]
    arrow_labels = ["API / Cloud", "GSM Cellular", "OS Wake"]
    for idx, ax in enumerate(arrow_xs):
        arr = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, y_top + Inches(1.1), Inches(0.5), Inches(0.26))
        arr.fill.solid()
        arr.fill.fore_color.rgb = C_EMERALD
        arr.line.fill.background()

        tb_a = s3.shapes.add_textbox(ax - Inches(0.25), y_top + Inches(0.78), Inches(1.0), Inches(0.25))
        tf_a = tb_a.text_frame
        tf_a.margin_left = tf_a.margin_top = tf_a.margin_right = tf_a.margin_bottom = 0
        p_a = tf_a.paragraphs[0]
        p_a.text = arrow_labels[idx]
        p_a.font.size = Pt(8.5)
        p_a.font.bold = True
        p_a.font.color.rgb = C_MUTED
        p_a.alignment = PP_ALIGN.CENTER

    loop_band = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.75), Inches(11.333), Inches(1.7))
    loop_band.fill.solid()
    loop_band.fill.fore_color.rgb = C_CARD
    loop_band.line.color.rgb = C_EMERALD_BORDER
    loop_band.line.width = Pt(1.5)

    ib = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.95), Inches(3.2), Inches(1.3))
    ib.fill.solid()
    ib.fill.fore_color.rgb = C_EMERALD_SOFT
    ib.line.color.rgb = C_EMERALD_BORDER
    ib.line.width = Pt(1)
    tf_ib = ib.text_frame
    tf_ib.word_wrap = True
    p = tf_ib.paragraphs[0]
    p.text = "CLOSED-LOOP FEEDBACK"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT
    p.alignment = PP_ALIGN.CENTER
    p2 = tf_ib.add_paragraph()
    p2.text = "\nReal-Time Command Telemetry\n100% Troop Accountability"
    p2.font.size = Pt(10)
    p2.font.color.rgb = C_TEXT_BODY
    p2.alignment = PP_ALIGN.CENTER

    tb_loop = s3.shapes.add_textbox(Inches(4.7), Inches(5.0), Inches(7.3), Inches(1.2))
    tf_loop = tb_loop.text_frame
    tf_loop.word_wrap = True
    p = tf_loop.paragraphs[0]
    p.text = "◄── Returning Telemetry: Instant Verification to HQ Dashboard"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT

    p2 = tf_loop.add_paragraph()
    p2.text = "When the soldier taps '✓ I RECEIVED THIS ORDER', the siren kills immediately and confirmation is transmitted back. The Headquarters Web Dashboard updates live in real time with individual soldier response timestamps, CDC battalion tallies, and true mobilization readiness."
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_TEXT_BODY

    # ==========================================
    # SLIDE 4: THE 4 CORE CAPABILITIES
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Core Capabilities", "The 4 Pillars of the Mobile Engine", 4, 7)

    tech_points = [
        ("1", "Hardware Alarm Override", "Bypasses Silent, Mute, and Do Not Disturb by using privileged Android USAGE_ALARM stream at maximum hardware volume."),
        ("2", "Lockscreen Takeover", "Forces the phone display ON even while locked, presenting the official 10RCDG emblem and mobilization order."),
        ("3", "100% Offline Capability", "Triggered over basic GSM text message payload. Does NOT require mobile data, Wi-Fi, or internet connection on the soldier's phone."),
        ("4", "Instant Acknowledgment", "Soldier taps '✓ I Received This Order' to silence the siren and confirm their mobilization status to Headquarters.")
    ]

    top = Inches(1.9)
    for num, title, desc in tech_points:
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top, Inches(11.333), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        num_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), top + Inches(0.18), Inches(0.6), Inches(0.6))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = C_EMERALD_SOFT
        num_box.line.color.rgb = C_EMERALD_BORDER
        num_box.line.width = Pt(1)
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(14)
        p_n.font.bold = True
        p_n.font.color.rgb = C_EMERALD_TEXT
        p_n.alignment = PP_ALIGN.CENTER

        tb = s4.shapes.add_textbox(Inches(2.1), top + Inches(0.15), Inches(3.2), Inches(0.65))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        tb_d = s4.shapes.add_textbox(Inches(5.4), top + Inches(0.15), Inches(6.7), Inches(0.65))
        tf_d = tb_d.text_frame
        tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = C_TEXT_BODY

        top += Inches(1.15)

    # ==========================================
    # SLIDE 5: LIVE DEMONSTRATION WORKFLOW
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Live Demonstration", "Hands-On Test: Web Dispatch to Soldier Siren", 5, 7)

    steps = [
        ("STEP 1", "Officer Dispatches Alert", "Commander selects target unit (e.g. 1001st CDC) on the web portal and clicks 'TRANSMIT EMERGENCY SIREN'.", C_EMERALD, C_EMERALD_SOFT),
        ("STEP 2", "Phone Lights Up & Sounds Siren", "The soldier's muted, locked phone immediately turns on, displays the 10RCDG emblem, and sounds a loud tactical siren.", C_AMBER, C_AMBER_SOFT),
        ("STEP 3", "Troop Acknowledges Order", "Soldier taps '✓ I RECEIVED THIS ORDER'. The alarm silences and compliance is recorded for headquarters.", C_EMERALD, C_EMERALD_SOFT)
    ]

    left = Inches(1.0)
    for step_num, title, desc, accent, bg_pill in steps:
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.5), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        badge = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.4), Inches(2.4), Inches(2.7), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_pill
        badge.line.color.rgb = accent
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = step_num
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = accent
        p_b.alignment = PP_ALIGN.CENTER

        tb = s5.shapes.add_textbox(left + Inches(0.35), Inches(3.3), Inches(2.8), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_BODY

        left += Inches(3.9)

    # ==========================================
    # SLIDE 6: SENDER ID REQUIREMENT
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Strategic Requirement", "The Goal: Securing an Official Alphanumeric SENDER ID", 6, 7)

    comp = [
        ("CURRENT PROTOTYPE (SIM Gateway)", [
            "Sent from an ordinary 11-digit number (0917-xxx-xxxx)",
            "Telco firewalls block all web links (.com, .app) due to NTC anti-smishing laws",
            "Slow speed: limited to 10–15 SMS per minute on SIM cards"
        ], C_RED, C_RED_SOFT),
        ("WITH OFFICIAL SENDER ID ('10RCDG')", [
            "Displays official military sender name: '10RCDG' or 'RESCOM-PA'",
            "100% Whitelisted by Globe & Smart: Zero spam blocking, allows official download links",
            "High-speed broadcast: Dispatches 100+ SMS per second to thousands of troops instantly"
        ], C_EMERALD, C_EMERALD_SOFT)
    ]

    left = Inches(1.0)
    for title, bullets, accent, bg_pill in comp:
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(5.45), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        h_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.4), Inches(2.3), Inches(4.65), Inches(0.6))
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = bg_pill
        h_box.line.color.rgb = accent
        h_box.line.width = Pt(1)
        tf_h = h_box.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = title
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = accent
        p_h.alignment = PP_ALIGN.CENTER

        tb = s6.shapes.add_textbox(left + Inches(0.4), Inches(3.1), Inches(4.65), Inches(2.9))
        tf = tb.text_frame
        tf.word_wrap = True

        for idx, b in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            prefix = "✗  " if accent == C_RED else "✓  "
            p.text = prefix + b + "\n"
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT if accent != C_RED else C_RED

        left += Inches(5.88)

    # ==========================================
    # SLIDE 7: LIVE DOWNLOAD PORTAL & QR CODE (NEW REPLACEMENT!)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    if os.path.exists(bg_slide_img):
        s7.shapes.add_picture(bg_slide_img, 0, 0, width=Inches(13.333), height=Inches(7.5))
    else:
        set_bg(s7)

    # Main Card
    card7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.85), Inches(11.333), Inches(5.8))
    card7.fill.solid()
    card7.fill.fore_color.rgb = C_CARD
    card7.line.color.rgb = C_BORDER
    card7.line.width = Pt(1.5)

    # LEFT COLUMN: Instructions & Live Demo Action
    if os.path.exists(logo_path):
        s7.shapes.add_picture(logo_path, Inches(1.4), Inches(1.15), width=Inches(0.95))
    if os.path.exists(seal_path):
        s7.shapes.add_picture(seal_path, Inches(2.5), Inches(1.15), width=Inches(0.95))

    badge7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.6), Inches(1.2), Inches(3.8), Inches(0.36))
    badge7.fill.solid()
    badge7.fill.fore_color.rgb = C_EMERALD_SOFT
    badge7.line.color.rgb = C_EMERALD_BORDER
    badge7.line.width = Pt(1)
    tf = badge7.text_frame
    p = tf.paragraphs[0]
    p.text = "LIVE HANDS-ON DEMONSTRATION"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT
    p.alignment = PP_ALIGN.CENTER

    tb_left = s7.shapes.add_textbox(Inches(1.4), Inches(2.2), Inches(6.0), Inches(3.2))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "Install & Test on Your Phone Now"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_TEXT

    p2 = tf_l.add_paragraph()
    p2.text = "Official Public Mobile Download Portal:"
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = C_EMERALD_TEXT

    p_url = tf_l.add_paragraph()
    p_url.text = "👉 https://rescom-alert.vercel.app/download"
    p_url.font.size = Pt(13)
    p_url.font.bold = True
    p_url.font.color.rgb = C_BLUE

    steps_l = [
        "1. Open your phone camera and scan the QR code.",
        "2. Tap 'DOWNLOAD SOLDIER APP' (No login required).",
        "3. If prompted, tap 'Download anyway' (Official military APK).",
        "4. Open app & tap 'Enable' for Full-Screen Siren permissions."
    ]
    for st in steps_l:
        p = tf_l.add_paragraph()
        p.text = st
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_BODY

    # Live Telemetry Notice Callout Box
    live_callout = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(5.1), Inches(6.0), Inches(0.85))
    live_callout.fill.solid()
    live_callout.fill.fore_color.rgb = C_AMBER_SOFT
    live_callout.line.color.rgb = C_AMBER_BORDER
    live_callout.line.width = Pt(1)
    tf_lc = live_callout.text_frame
    tf_lc.word_wrap = True
    p = tf_lc.paragraphs[0]
    p.text = "⚡ LIVE DEDUPLICATION COUNTER DEMO:"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_AMBER
    p2 = tf_lc.add_paragraph()
    p2.text = "Every officer who installs will appear as +1 Unique Device on the Commander Dashboard in real time!"
    p2.font.size = Pt(10)
    p2.font.color.rgb = C_TEXT

    # RIGHT COLUMN: Clean QR Code Card
    qr_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.2), Inches(3.9), Inches(5.0))
    qr_card.fill.solid()
    qr_card.fill.fore_color.rgb = C_CARD
    qr_card.line.color.rgb = C_EMERALD_BORDER
    qr_card.line.width = Pt(1.5)

    # Embedded Emerald QR Image
    if os.path.exists(qr_img):
        s7.shapes.add_picture(qr_img, Inches(8.35), Inches(1.5), width=Inches(3.0))

    # QR Subtext Pill
    qr_pill = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(4.7), Inches(3.2), Inches(0.4))
    qr_pill.fill.solid()
    qr_pill.fill.fore_color.rgb = C_EMERALD_SOFT
    qr_pill.line.color.rgb = C_EMERALD_BORDER
    qr_pill.line.width = Pt(1)
    tf_qp = qr_pill.text_frame
    p = tf_qp.paragraphs[0]
    p.text = "SCAN WITH ANDROID CAMERA"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT
    p.alignment = PP_ALIGN.CENTER

    tb_qr_sub = s7.shapes.add_textbox(Inches(8.1), Inches(5.2), Inches(3.5), Inches(0.8))
    tf_qs = tb_qr_sub.text_frame
    tf_qs.word_wrap = True
    p = tf_qs.paragraphs[0]
    p.text = "Compatible with Android 7.0 to Android 15+"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED
    p.alignment = PP_ALIGN.CENTER

    p2 = tf_qs.add_paragraph()
    p2.text = "Javier Siliacay • USTP Autotronics"
    p2.font.size = Pt(10)
    p2.font.bold = True
    p2.font.color.rgb = C_TEXT
    p2.alignment = PP_ALIGN.CENTER

    s7.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "Colonel, Sir/Ma'am, to conclude this presentation:\n"
        "I invite all officers in this room to point your phone cameras at this QR code right now.\n"
        "It will take you directly to our live military download portal at https://rescom-alert.vercel.app/download.\n"
        "As you tap download, watch my laptop screen — our live system will register each unique device on the Commander Dashboard in real time.\n"
        "Thank you very much. The floor is now open for questions and live testing."
    )

    out_file = os.path.abspath("10RCDG_RESCOM_ALERT_Final.pptx")
    prs.save(out_file)
    print(f"Final presentation saved at: {out_file}")

    public_out = os.path.abspath("public/10RCDG_RESCOM_ALERT_Final.pptx")
    try:
        prs.save(public_out)
    except Exception:
        pass

if __name__ == "__main__":
    create_deck_with_download_slide()
