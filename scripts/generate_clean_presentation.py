import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_clean_tech_deck_with_diagram():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Clean System Tokens
    C_BG = RGBColor(248, 250, 252)          # slate-50
    C_CARD = RGBColor(255, 255, 255)        # pure white
    C_BORDER = RGBColor(226, 232, 240)      # slate-200
    C_TEXT = RGBColor(15, 23, 42)           # slate-900
    C_TEXT_BODY = RGBColor(71, 85, 105)     # slate-600
    C_MUTED = RGBColor(148, 163, 184)       # slate-400

    # Accents
    C_EMERALD = RGBColor(6, 95, 70)         # emerald-800
    C_EMERALD_TEXT = RGBColor(4, 120, 87)   # emerald-700
    C_EMERALD_SOFT = RGBColor(236, 253, 245)# emerald-50
    C_EMERALD_BORDER = RGBColor(167, 243, 208)

    C_BLUE = RGBColor(29, 78, 216)          # blue-700
    C_BLUE_SOFT = RGBColor(239, 246, 255)   # blue-50
    C_BLUE_BORDER = RGBColor(191, 219, 254)

    C_AMBER = RGBColor(217, 119, 6)         # amber-600
    C_AMBER_SOFT = RGBColor(255, 251, 235)  # amber-50
    C_RED = RGBColor(220, 38, 38)           # red-600
    C_RED_SOFT = RGBColor(254, 242, 242)    # red-50

    logo_path = os.path.abspath("public/rescom-emblem.jpg")
    seal_path = os.path.abspath("public/rescom-pa-seal.png")

    blank_layout = prs.slide_layouts[6]

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, tag, title, slide_num, total=7):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.5), Inches(2.8), Inches(0.36))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_EMERALD_SOFT
        badge.line.color.rgb = C_EMERALD_BORDER
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = tag.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = C_EMERALD_TEXT
        p_b.alignment = PP_ALIGN.CENTER

        tb = slide.shapes.add_textbox(Inches(1.0), Inches(0.95), Inches(9.5), Inches(0.65))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        if os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(11.6), Inches(0.5), width=Inches(0.75))

        tb_f = slide.shapes.add_textbox(Inches(1.0), Inches(6.9), Inches(8.0), Inches(0.3))
        tf_f = tb_f.text_frame
        tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
        p = tf_f.paragraphs[0]
        p.text = "10RCDG RESCOM-ALERT  •  Technical Briefing & Command Defense"
        p.font.size = Pt(10)
        p.font.color.rgb = C_MUTED

        tb_n = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.0), Inches(0.3))
        tf_n = tb_n.text_frame
        tf_n.margin_left = tf_n.margin_top = tf_n.margin_right = tf_n.margin_bottom = 0
        p = tf_n.paragraphs[0]
        p.text = f"{slide_num:02d} / {total:02d}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # SLIDE 1: CLEAN TITLE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_CARD
    card1.line.color.rgb = C_BORDER
    card1.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(5.5), Inches(1.4), width=Inches(1.1))
    if os.path.exists(seal_path):
        s1.shapes.add_picture(seal_path, Inches(6.8), Inches(1.4), width=Inches(1.1))

    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(2.7), Inches(4.7), Inches(0.38))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = C_EMERALD_SOFT
    badge1.line.color.rgb = C_EMERALD_BORDER
    badge1.line.width = Pt(1)
    tf = badge1.text_frame
    p = tf.paragraphs[0]
    p.text = "10RCDG RAPID MOBILIZATION SYSTEM"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT
    p.alignment = PP_ALIGN.CENTER

    tb = s1.shapes.add_textbox(Inches(2.0), Inches(3.3), Inches(9.333), Inches(1.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "RESCOM-ALERT"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_TEXT
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "Emergency Siren Override & Instant Troop Recall Technology"
    p.font.size = Pt(18)
    p.font.color.rgb = C_TEXT_BODY
    p.alignment = PP_ALIGN.CENTER

    tb_sub = s1.shapes.add_textbox(Inches(2.0), Inches(5.2), Inches(9.333), Inches(0.8))
    tf_sub = tb_sub.text_frame
    p = tf_sub.paragraphs[0]
    p.text = "Javier Siliacay (aka vier from USTP Autotronics)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT
    p.alignment = PP_ALIGN.CENTER

    p2 = tf_sub.add_paragraph()
    p2.text = "Lead Systems Developer  •  October 2026"
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_MUTED
    p2.alignment = PP_ALIGN.CENTER

    s1.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "Good day, Colonel, and members of the 10RCDG Command Staff.\n"
        "Today I am presenting RESCOM-ALERT — a purpose-built tactical technology designed to solve one critical problem: ensuring emergency mobilization orders wake up and alert our reservists, even when their phones are locked, muted, or offline."
    )

    # ==========================================
    # SLIDE 2: THE PROBLEM (Simple & High Impact)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "The Problem", "Why Conventional Messaging Fails in Emergencies", 2, 7)

    problems = [
        ("74% SILENT", "Phones on Mute & DND", "Troops sleep through or miss standard SMS dings and chat notifications during Red Alerts.", C_RED, C_RED_SOFT),
        ("NO DATA", "Chat Apps Fail in Storms", "Messenger and Viber require active internet. When cell towers lose data, chat groups fail.", C_AMBER, C_AMBER_SOFT),
        ("NO HEADCOUNT", "Zero Acknowledgment", "Commanders have no way to verify who received the order and who is actively en route.", C_TEXT, RGBColor(241, 245, 249))
    ]

    left = Inches(1.0)
    for stat, header, desc, accent, bg_pill in problems:
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.5), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        b_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.4), Inches(2.4), Inches(2.7), Inches(0.8))
        b_box.fill.solid()
        b_box.fill.fore_color.rgb = bg_pill
        b_box.line.color.rgb = accent
        b_box.line.width = Pt(1)
        tf_b = b_box.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = stat
        p_b.font.size = Pt(20)
        p_b.font.bold = True
        p_b.font.color.rgb = accent
        p_b.alignment = PP_ALIGN.CENTER

        tb = s2.shapes.add_textbox(left + Inches(0.35), Inches(3.5), Inches(2.8), Inches(2.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = header
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_BODY

        left += Inches(3.9)

    s2.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "Sir, our current communication has three major fatal flaws during rapid mobilization:\n"
        "1. Over 70% of personnel keep phones on silent or Do Not Disturb. Regular text dings don't wake people up.\n"
        "2. Chat apps like Messenger depend entirely on internet data. When typhoons or earthquakes knock out cell data, communication is dead.\n"
        "3. Commanders have zero headcount visibility."
    )

    # ==========================================
    # SLIDE 3: SYSTEM ARCHITECTURE DIAGRAM (NEW!)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "System Architecture", "End-to-End Operational Signal Flow", 3, 7)

    # Block 1: Headquarters Command
    b1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(3.2), Inches(4.3))
    b1.fill.solid()
    b1.fill.fore_color.rgb = C_CARD
    b1.line.color.rgb = C_BORDER
    b1.line.width = Pt(1.5)

    badge_b1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(2.3), Inches(2.6), Inches(0.5))
    badge_b1.fill.solid()
    badge_b1.fill.fore_color.rgb = C_EMERALD_SOFT
    badge_b1.line.color.rgb = C_EMERALD_BORDER
    badge_b1.line.width = Pt(1)
    tf = badge_b1.text_frame
    p = tf.paragraphs[0]
    p.text = "10RCDG HEADQUARTERS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_TEXT
    p.alignment = PP_ALIGN.CENTER

    tb_b1 = s3.shapes.add_textbox(Inches(1.2), Inches(3.0), Inches(2.8), Inches(3.1))
    tf1 = tb_b1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Command Web Portal\n"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_TEXT

    b1_items = [
        "Officer logs into secure console",
        "Selects Target CDC (e.g. 1001st)",
        "Enters emergency directive",
        "Clicks 'Transmit Siren Alert'"
    ]
    for it in b1_items:
        p = tf1.add_paragraph()
        p.text = "•  " + it
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_BODY

    # Arrow 1 -> 2
    arr1 = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.35), Inches(3.8), Inches(0.6), Inches(0.35))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = C_EMERALD
    arr1.line.fill.background()

    # Block 2: Telco Gateway / SENDER ID
    b2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(2.0), Inches(3.2), Inches(4.3))
    b2.fill.solid()
    b2.fill.fore_color.rgb = C_CARD
    b2.line.color.rgb = C_BORDER
    b2.line.width = Pt(1.5)

    badge_b2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.4), Inches(2.3), Inches(2.6), Inches(0.5))
    badge_b2.fill.solid()
    badge_b2.fill.fore_color.rgb = C_BLUE_SOFT
    badge_b2.line.color.rgb = C_BLUE_BORDER
    badge_b2.line.width = Pt(1)
    tf = badge_b2.text_frame
    p = tf.paragraphs[0]
    p.text = "TELCO INFRASTRUCTURE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.alignment = PP_ALIGN.CENTER

    tb_b2 = s3.shapes.add_textbox(Inches(5.3), Inches(3.0), Inches(2.8), Inches(3.1))
    tf2 = tb_b2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "National GSM SMS Airwave\n"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_TEXT

    b2_items = [
        "Smart / Globe cellular tower network",
        "Operates via basic 2G/3G/4G/5G SMS",
        "ZERO mobile data or Wi-Fi required",
        "Delivers directly to target SIMs in 2s"
    ]
    for it in b2_items:
        p = tf2.add_paragraph()
        p.text = "•  " + it
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_BODY

    # Arrow 2 -> 3
    arr2 = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.45), Inches(3.8), Inches(0.6), Inches(0.35))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = C_EMERALD
    arr2.line.fill.background()

    # Block 3: Soldier Device Hardware Override
    b3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(2.0), Inches(3.2), Inches(4.3))
    b3.fill.solid()
    b3.fill.fore_color.rgb = C_CARD
    b3.line.color.rgb = C_BORDER
    b3.line.width = Pt(1.5)

    badge_b3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(2.3), Inches(2.6), Inches(0.5))
    badge_b3.fill.solid()
    badge_b3.fill.fore_color.rgb = C_AMBER_SOFT
    badge_b3.line.color.rgb = C_AMBER
    badge_b3.line.width = Pt(1)
    tf = badge_b3.text_frame
    p = tf.paragraphs[0]
    p.text = "SOLDIER FIELD DEVICE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_AMBER
    p.alignment = PP_ALIGN.CENTER

    tb_b3 = s3.shapes.add_textbox(Inches(9.4), Inches(3.0), Inches(2.8), Inches(3.1))
    tf3 = tb_b3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Native Android Override\n"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_TEXT

    b3_items = [
        "SmsAlertReceiver intercepts text",
        "USAGE_ALARM breaks silent/mute",
        "Wakes screen with 10RCDG emblem",
        "Troop taps 'I Received Order'"
    ]
    for it in b3_items:
        p = tf3.add_paragraph()
        p.text = "•  " + it
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_BODY

    s3.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "Sir, this diagram shows the entire signal pipeline in 3 clear steps:\n"
        "Step 1: Headquarters initiates the order on the Command Web Portal.\n"
        "Step 2: The alert travels over the standard cellular airwaves via basic SMS — notice there is ZERO reliance on internet data. Even if mobile data is dead, basic SMS penetrates.\n"
        "Step 3: The soldier's Android phone receives the payload. The background receiver intercepts it, overrides all volume settings, turns on the locked screen with the 10RCDG emblem, and sounds the siren until acknowledged."
    )

    # ==========================================
    # SLIDE 4: THE 4 CORE CAPABILITIES
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Core Capabilities", "The 4 Pillars of the Mobile Engine", 4, 7)

    tech_points = [
        ("1", "Hardware Alarm Override", "Bypasses Silent, Mute, and Do Not Disturb by using privileged Android USAGE_ALARM stream at maximum hardware volume."),
        ("2", "Lockscreen Takeover", "Forces the phone display ON even while locked, presenting the official 10RCDG emblem and mobilization order."),
        ("3", "100% Offline Capability", "Triggered over basic GSM text message payload. Does NOT require mobile data, Wi-Fi, or internet connection on the soldier's phone."),
        ("4", "Instant Acknowledgment", "Soldier taps '✓ I Received This Order' to silence the siren and confirm their mobilization status to Headquarters.")
    ]

    top = Inches(1.9)
    for num, title, desc in tech_points:
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top, Inches(11.333), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        num_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), top + Inches(0.18), Inches(0.6), Inches(0.6))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = C_EMERALD_SOFT
        num_box.line.color.rgb = C_EMERALD_BORDER
        num_box.line.width = Pt(1)
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(14)
        p_n.font.bold = True
        p_n.font.color.rgb = C_EMERALD_TEXT
        p_n.alignment = PP_ALIGN.CENTER

        tb = s4.shapes.add_textbox(Inches(2.1), top + Inches(0.15), Inches(3.2), Inches(0.65))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        tb_d = s4.shapes.add_textbox(Inches(5.4), top + Inches(0.15), Inches(6.7), Inches(0.65))
        tf_d = tb_d.text_frame
        tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = C_TEXT_BODY

        top += Inches(1.15)

    s4.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "To summarize the four core capabilities of the mobile engine:\n"
        "First: Volume slider bypass.\n"
        "Second: Screen wake takeover.\n"
        "Third: Total offline operation.\n"
        "Fourth: Acknowledgment accountability."
    )

    # ==========================================
    # SLIDE 5: LIVE DEMONSTRATION WORKFLOW
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Live Demonstration", "Hands-On Test: Web Dispatch to Soldier Siren", 5, 7)

    steps = [
        ("STEP 1", "Officer Dispatches Alert", "Commander selects target unit (e.g. 1001st CDC) on the web portal and clicks 'TRANSMIT EMERGENCY SIREN'.", C_EMERALD, C_EMERALD_SOFT),
        ("STEP 2", "Phone Lights Up & Sounds Siren", "The soldier's muted, locked phone immediately turns on, displays the 10RCDG emblem, and sounds a loud tactical siren.", C_AMBER, C_AMBER_SOFT),
        ("STEP 3", "Troop Acknowledges Order", "Soldier taps '✓ I RECEIVED THIS ORDER'. The alarm silences and compliance is recorded for headquarters.", C_EMERALD, C_EMERALD_SOFT)
    ]

    left = Inches(1.0)
    for step_num, title, desc, accent, bg_pill in steps:
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.5), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        badge = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.4), Inches(2.4), Inches(2.7), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_pill
        badge.line.color.rgb = accent
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = step_num
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = accent
        p_b.alignment = PP_ALIGN.CENTER

        tb = s5.shapes.add_textbox(left + Inches(0.35), Inches(3.3), Inches(2.8), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_BODY

        left += Inches(3.9)

    s5.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "[PAUSE FOR LIVE DEMONSTRATION]\n"
        "Sir, to demonstrate this technology, I have placed a test phone right here on the table — it is locked and set to silent.\n"
        "From my laptop portal, I am now firing a live drill alert...\n"
        "[TRIGGER ALERT]\n"
        "As you can hear and see, the phone immediately activates at full volume. The order cannot be missed."
    )

    # ==========================================
    # SLIDE 6: SENDER ID REQUIREMENT
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Strategic Requirement", "The Goal: Securing an Official Alphanumeric SENDER ID", 6, 7)

    comp = [
        ("CURRENT PROTOTYPE (SIM Gateway)", [
            "Sent from an ordinary 11-digit number (0917-xxx-xxxx)",
            "Telco firewalls block all web links (.com, .app) due to NTC anti-smishing laws",
            "Slow speed: limited to 10–15 SMS per minute on SIM cards"
        ], C_RED, C_RED_SOFT),
        ("WITH OFFICIAL SENDER ID ('10RCDG')", [
            "Displays official military sender name: '10RCDG' or 'RESCOM-PA'",
            "100% Whitelisted by Globe & Smart: Zero spam blocking, allows official download links",
            "High-speed broadcast: Dispatches 100+ SMS per second to thousands of troops instantly"
        ], C_EMERALD, C_EMERALD_SOFT)
    ]

    left = Inches(1.0)
    for title, bullets, accent, bg_pill in comp:
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(5.45), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        h_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.4), Inches(2.3), Inches(4.65), Inches(0.6))
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = bg_pill
        h_box.line.color.rgb = accent
        h_box.line.width = Pt(1)
        tf_h = h_box.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = title
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = accent
        p_h.alignment = PP_ALIGN.CENTER

        tb = s6.shapes.add_textbox(left + Inches(0.4), Inches(3.1), Inches(4.65), Inches(2.9))
        tf = tb.text_frame
        tf.word_wrap = True

        for idx, b in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            prefix = "✗  " if accent == C_RED else "✓  "
            p.text = prefix + b + "\n"
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT if accent != C_RED else C_RED

        left += Inches(5.88)

    s6.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "Sir, this is why we are here today.\n"
        "The software and technology are already 100% built and functioning.\n"
        "What we need from national command is budget allocation for an Official Alphanumeric SENDER ID (such as '10RCDG').\n"
        "With a verified Sender ID, our alerts arrive under the official 10RCDG banner, bypass carrier spam blocking, and dispatch to thousands of troops in seconds."
    )

    # ==========================================
    # SLIDE 7: CONCLUSION & THE ASK
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)

    card7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
    card7.fill.solid()
    card7.fill.fore_color.rgb = C_CARD
    card7.line.color.rgb = C_BORDER
    card7.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s7.shapes.add_picture(logo_path, Inches(6.0), Inches(1.3), width=Inches(1.3))

    tbox7 = s7.shapes.add_textbox(Inches(2.0), Inches(2.8), Inches(9.333), Inches(2.0))
    tf7 = tbox7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "“When Every Second Counts, Silence is Not an Option.”"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_TEXT
    p.alignment = PP_ALIGN.CENTER

    p2 = tf7.add_paragraph()
    p2.text = "\nThe software is ready. We respectfully request command endorsement\nand budget allocation for the official 10RCDG SENDER ID."
    p2.font.size = Pt(16)
    p2.font.color.rgb = C_EMERALD_TEXT
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf7.add_paragraph()
    p3.text = "\nFloor is open for Questions & Live Testing."
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = C_MUTED
    p3.alignment = PP_ALIGN.CENTER

    tb_sub7 = s7.shapes.add_textbox(Inches(2.0), Inches(5.3), Inches(9.333), Inches(0.8))
    tf_sub7 = tb_sub7.text_frame
    p = tf_sub7.paragraphs[0]
    p.text = "Javier Siliacay (aka vier from USTP Autotronics)  •  Lead Developer"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_TEXT
    p.alignment = PP_ALIGN.CENTER

    s7.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT:\n"
        "In conclusion, Sir/Ma'am:\n"
        "The technology works. The app is ready. The command console is operational.\n"
        "With your endorsement and budget for the SENDER ID, 10RCDG will become the fastest, most technologically responsive mobilization command in the Philippine Army.\n"
        "Thank you, Sir. I am ready for your questions."
    )

    clean_path = os.path.abspath("10RCDG_RESCOM_ALERT_Clean.pptx")
    try:
        prs.save(clean_path)
        print(f"Clean deck successfully generated at: {clean_path}")
    except Exception as e:
        print(f"Could not overwrite {clean_path}: {e}")

    diagram_path = os.path.abspath("10RCDG_RESCOM_ALERT_Architecture.pptx")
    prs.save(diagram_path)
    print(f"Deck with diagram saved at: {diagram_path}")

    public_diagram = os.path.abspath("public/10RCDG_RESCOM_ALERT_Architecture.pptx")
    try:
        prs.save(public_diagram)
    except Exception:
        pass

if __name__ == "__main__":
    create_clean_tech_deck_with_diagram()
